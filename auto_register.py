# -*- coding: utf-8 -*-
import random
import string
import os
import time
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from playwright.sync_api import sync_playwright
import openpyxl
from openpyxl import Workbook

EXCEL_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "register_data.xlsx")
VALIDATION_EXCEL_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "validation_results.xlsx")

FIRST_NAMES = ["张", "王", "李", "赵", "刘", "陈", "杨", "黄", "周", "吴", "徐", "孙", "马", "朱", "胡", "郭", "何", "高", "林", "罗"]
LAST_NAMES = ["伟", "芳", "娜", "秀英", "敏", "静", "丽", "强", "磊", "军", "洋", "勇", "艳", "杰", "娟", "涛", "明", "超", "秀兰", "霞"]

excel_lock = threading.Lock()
validation_data_list = []
validation_list_lock = threading.Lock()

def generate_random_string(length=8):
    return ''.join(random.choices(string.ascii_lowercase + string.digits, k=length))

def generate_random_email():
    username = generate_random_string(10)
    return f"{username}@163.com"

def generate_random_password():
    return generate_random_string(12) + random.choice(string.ascii_uppercase) + random.choice(string.digits)

def generate_random_name():
    first_name = random.choice(FIRST_NAMES)
    last_name = random.choice(LAST_NAMES)
    return first_name + last_name

def generate_random_age():
    return str(random.randint(18, 60))

def generate_random_phone():
    prefixes = ["130", "131", "132", "133", "134", "135", "136", "137", "138", "139",
                "150", "151", "152", "153", "155", "156", "157", "158", "159",
                "170", "176", "177", "178",
                "180", "181", "182", "183", "184", "185", "186", "187", "188", "189"]
    prefix = random.choice(prefixes)
    suffix = ''.join(random.choices(string.digits, k=8))
    return prefix + suffix

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "注册数据"
        headers = ["序号", "用户名", "密码", "邮箱", "姓名", "年龄", "手机号", "注册时间", "注册状态", "登录状态"]
        ws.append(headers)
        wb.save(EXCEL_FILE)
        print(f"创建Excel文件: {EXCEL_FILE}")
    return EXCEL_FILE

def init_validation_excel():
    if not os.path.exists(VALIDATION_EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "字段校验结果"
        headers = ["序号", "账号", "修改字段", "原字段值", "目标修改值", "测试类型", "接口返回结果", "校验是否通过", "测试时间"]
        ws.append(headers)
        wb.save(VALIDATION_EXCEL_FILE)
        print(f"创建校验Excel文件: {VALIDATION_EXCEL_FILE}")
    return VALIDATION_EXCEL_FILE

def save_to_excel(data):
    with excel_lock:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        ws.append(data)
        wb.save(EXCEL_FILE)

def add_validation_data(data):
    with validation_list_lock:
        validation_data_list.append(data)

def save_all_validation_data():
    if not validation_data_list:
        return
    with excel_lock:
        wb = openpyxl.load_workbook(VALIDATION_EXCEL_FILE)
        ws = wb.active
        for data in validation_data_list:
            ws.append(data)
        wb.save(VALIDATION_EXCEL_FILE)

def perform_register(page, user_data, task_id):
    username = user_data["username"]
    password = user_data["password"]
    email = user_data["email"]
    name = user_data["name"]
    age = user_data["age"]
    phone = user_data["phone"]
    
    register_start_time = datetime.now()
    
    print(f"[任务{task_id}] 填写注册信息...")
    
    all_inputs = page.query_selector_all('input')
    for inp in all_inputs:
        try:
            placeholder = inp.get_attribute('placeholder') or ""
            input_type = inp.get_attribute('type') or ""
            
            if '账号' in placeholder:
                inp.fill(username)
            elif '姓名' in placeholder:
                inp.fill(name)
            elif '年龄' in placeholder:
                inp.fill(age)
            elif '密码' in placeholder and input_type == 'password':
                inp.fill(password)
            elif '手机' in placeholder:
                inp.fill(phone)
            elif '邮箱' in placeholder:
                inp.fill(email)
        except:
            pass
    
    page.wait_for_timeout(500)
    
    register_btn = page.query_selector('button:has-text("注册")')
    
    register_status = "失败"
    if register_btn:
        print(f"[任务{task_id}] 点击注册按钮...")
        register_btn.click()
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(2000)
        register_status = "成功"
    
    return {
        "register_time": register_start_time.strftime("%Y-%m-%d %H:%M:%S"),
        "register_status": register_status
    }

def perform_login(page, user_data, task_id):
    username = user_data["username"]
    password = user_data["password"]
    
    print(f"[任务{task_id}] 跳转到登录页面...")
    page.goto("http://39.107.109.8:8082/login", timeout=30000)
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(1000)
    
    print(f"[任务{task_id}] 填写登录信息...")
    
    all_inputs = page.query_selector_all('input')
    for inp in all_inputs:
        try:
            placeholder = inp.get_attribute('placeholder') or ""
            input_type = inp.get_attribute('type') or ""
            
            if '账号' in placeholder:
                inp.fill(username)
            elif '密码' in placeholder and input_type == 'password':
                inp.fill(password)
        except:
            pass
    
    page.wait_for_timeout(500)
    
    login_btn = page.query_selector('button:has-text("登录")')
    
    login_status = "失败"
    
    if login_btn:
        print(f"[任务{task_id}] 点击登录按钮...")
        login_btn.click()
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(2000)
        
        if 'profile' in page.url:
            login_status = "成功"
    
    return {
        "login_status": login_status
    }

def generate_test_values(field_name, original_value):
    test_cases = []
    
    test_cases.append({
        "value": "",
        "test_type": "空字段"
    })
    
    test_cases.append({
        "value": "a" * 100,
        "test_type": "100字符长字段"
    })
    
    test_cases.append({
        "value": "12345",
        "test_type": "纯数字"
    })
    
    test_cases.append({
        "value": "abc123",
        "test_type": "字母数字混合"
    })
    
    test_cases.append({
        "value": "!@#$%^&*()_+-=[]{}|;':\",./<>?",
        "test_type": "特殊符号"
    })
    
    test_cases.append({
        "value": "<script>alert('xss')</script>",
        "test_type": "XSS攻击测试"
    })
    
    test_cases.append({
        "value": "'; DROP TABLE users; --",
        "test_type": "SQL注入测试"
    })
    
    if field_name in ["name", "姓名"]:
        test_cases.append({
            "value": "测试用户",
            "test_type": "正常中文"
        })
        test_cases.append({
            "value": "Test User",
            "test_type": "英文姓名"
        })
    
    if field_name in ["email", "邮箱"]:
        test_cases.append({
            "value": "invalid-email",
            "test_type": "无效邮箱格式"
        })
        test_cases.append({
            "value": "test@",
            "test_type": "不完整邮箱"
        })
    
    if field_name in ["phone", "手机", "手机号"]:
        test_cases.append({
            "value": "123",
            "test_type": "短手机号"
        })
        test_cases.append({
            "value": "abcdefghijk",
            "test_type": "非数字手机号"
        })
    
    if field_name in ["age", "年龄"]:
        test_cases.append({
            "value": "-1",
            "test_type": "负数年龄"
        })
        test_cases.append({
            "value": "200",
            "test_type": "超大年龄"
        })
        test_cases.append({
            "value": "abc",
            "test_type": "非数字年龄"
        })
    
    return test_cases

def enter_edit_mode(page, task_id):
    edit_btn = page.query_selector('button:has-text("编辑")')
    if edit_btn:
        edit_btn.click()
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(500)
        return True
    return False

def ensure_edit_mode(page, task_id):
    save_btn = page.query_selector('button:has-text("保存")')
    if not save_btn:
        enter_edit_mode(page, task_id)
        page.wait_for_timeout(300)

def find_field_input(page, field_name):
    all_inputs = page.query_selector_all('input')
    
    field_keywords = {
        "name": ["姓名"],
        "email": ["邮箱"],
        "phone": ["手机"],
        "age": ["年龄"]
    }
    
    keywords = field_keywords.get(field_name, [])
    
    for inp in all_inputs:
        try:
            placeholder = inp.get_attribute('placeholder') or ""
            for keyword in keywords:
                if keyword in placeholder:
                    return inp
        except:
            pass
    
    return None

def get_original_field_value(page, field_name):
    element = find_field_input(page, field_name)
    if element:
        try:
            return element.input_value()
        except:
            pass
    return ""

def test_field_validation(page, field_name, test_value, test_type, username, original_value, task_id, validation_index):
    ensure_edit_mode(page, task_id)
    
    element = find_field_input(page, field_name)
    
    if not element:
        print(f"[任务{task_id}] 未找到字段 {field_name} 的输入框")
        return {
            "success": False,
            "result": "未找到输入框"
        }
    
    try:
        element.fill("")
        page.wait_for_timeout(100)
        element.fill(str(test_value))
        page.wait_for_timeout(200)
        
        save_btn = page.query_selector('button:has-text("保存")')
        
        api_result = ""
        validation_passed = "否"
        
        if save_btn:
            save_btn.click()
            page.wait_for_timeout(1500)
            
            error_el = page.query_selector('.error-message')
            if error_el:
                try:
                    error_text = error_el.inner_text()
                    if error_text.strip():
                        api_result = f"错误提示: {error_text}"
                        validation_passed = "是(后端已拦截)"
                except:
                    pass
            
            if not api_result:
                success_el = page.query_selector('.success-message')
                if success_el:
                    try:
                        success_text = success_el.inner_text()
                        if success_text.strip():
                            api_result = f"成功: {success_text}"
                            validation_passed = "否(后端未拦截)"
                    except:
                        pass
            
            if not api_result:
                api_result = "请求已发送，无明确提示"
                validation_passed = "未知"
        else:
            api_result = "未找到保存按钮"
            validation_passed = "无法测试"
        
        validation_data = [
            validation_index,
            username,
            field_name,
            original_value,
            str(test_value)[:50],
            test_type,
            api_result[:100] if len(api_result) > 100 else api_result,
            validation_passed,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ]
        
        add_validation_data(validation_data)
        
        print(f"[任务{task_id}] 字段 {field_name} 测试类型 {test_type}: {validation_passed}")
        
        return {
            "success": True,
            "result": api_result,
            "validation_passed": validation_passed
        }
        
    except Exception as e:
        print(f"[任务{task_id}] 测试字段 {field_name} 时发生错误: {e}")
        return {
            "success": False,
            "result": str(e)
        }

def perform_field_validation(page, user_data, task_id, validation_start_index):
    print(f"[任务{task_id}] 开始字段校验测试...")
    
    if 'profile' not in page.url:
        print(f"[任务{task_id}] 不在个人信息页面，跳过字段校验")
        return []
    
    if not enter_edit_mode(page, task_id):
        print(f"[任务{task_id}] 无法进入编辑模式，跳过字段校验")
        return []
    
    fields_to_test = ["name", "email", "phone", "age"]
    field_display_names = {
        "name": "姓名",
        "email": "邮箱",
        "phone": "手机号",
        "age": "年龄"
    }
    
    all_results = []
    current_index = validation_start_index
    
    for field in fields_to_test:
        ensure_edit_mode(page, task_id)
        original_value = get_original_field_value(page, field)
        test_cases = generate_test_values(field_display_names.get(field, field), original_value)
        
        print(f"[任务{task_id}] 测试字段 {field_display_names.get(field, field)}, 原值: {original_value}")
        
        for test_case in test_cases:
            result = test_field_validation(
                page, 
                field, 
                test_case["value"], 
                test_case["test_type"],
                user_data["username"],
                original_value,
                task_id,
                current_index
            )
            all_results.append({
                "field": field,
                "test_type": test_case["test_type"],
                "test_value": test_case["value"],
                "original_value": original_value,
                "result": result
            })
            current_index += 1
            
            page.wait_for_timeout(300)
    
    return all_results

def single_task(task_id, user_data, validation_start_index):
    print(f"\n[任务{task_id}] 开始执行...")
    print(f"[任务{task_id}] 用户名: {user_data['username']}")
    
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context()
        page = context.new_page()
        
        try:
            print(f"[任务{task_id}] 正在打开网站...")
            page.goto("http://39.107.109.8:8082/", timeout=30000)
            page.wait_for_load_state("networkidle")
            page.wait_for_timeout(2000)
            
            print(f"[任务{task_id}] 查找注册链接...")
            register_link = page.query_selector('a:has-text("注册")')
            
            if register_link:
                print(f"[任务{task_id}] 点击注册链接...")
                register_link.click()
                page.wait_for_load_state("networkidle")
                page.wait_for_timeout(1000)
            
            register_result = perform_register(page, user_data, task_id)
            
            login_result = perform_login(page, user_data, task_id)
            
            validation_results = []
            if login_result["login_status"] == "成功":
                validation_results = perform_field_validation(page, user_data, task_id, validation_start_index)
            
            excel_data = [
                task_id,
                user_data["username"],
                user_data["password"],
                user_data["email"],
                user_data["name"],
                user_data["age"],
                user_data["phone"],
                register_result["register_time"],
                register_result["register_status"],
                login_result["login_status"]
            ]
            
            save_to_excel(excel_data)
            
            print(f"\n[任务{task_id}] ========== 执行完成 ==========")
            print(f"[任务{task_id}] 注册状态: {register_result['register_status']}")
            print(f"[任务{task_id}] 登录状态: {login_result['login_status']}")
            print(f"[任务{task_id}] 字段校验测试数: {len(validation_results)}")
            print(f"[任务{task_id}] ==============================\n")
            
            return {
                "task_id": task_id,
                "status": "成功",
                "register_status": register_result["register_status"],
                "login_status": login_result["login_status"],
                "validation_count": len(validation_results)
            }
            
        except Exception as e:
            print(f"[任务{task_id}] 发生错误: {e}")
            
            return {
                "task_id": task_id,
                "status": "失败",
                "error": str(e)
            }
        finally:
            browser.close()

def generate_user_data():
    return {
        "username": "user_" + generate_random_string(6),
        "password": generate_random_password(),
        "email": generate_random_email(),
        "name": generate_random_name(),
        "age": generate_random_age(),
        "phone": generate_random_phone()
    }

def run_parallel_register(num_tasks=5):
    global validation_data_list
    validation_data_list = []
    
    init_excel()
    init_validation_excel()
    
    print("=" * 60)
    print(f"开始并行执行 {num_tasks} 个注册任务")
    print("=" * 60)
    
    overall_start_time = datetime.now()
    
    users_data = [generate_user_data() for _ in range(num_tasks)]
    
    print("\n生成的用户信息:")
    for i, user in enumerate(users_data, 1):
        print(f"  任务{i}: {user['username']}")
    
    results = []
    validation_start_index = 1
    
    with ThreadPoolExecutor(max_workers=num_tasks) as executor:
        futures = {}
        for i, user in enumerate(users_data):
            future = executor.submit(single_task, i+1, user, validation_start_index)
            futures[future] = i+1
            validation_start_index += 50
        
        for future in as_completed(futures):
            task_id = futures[future]
            try:
                result = future.result()
                results.append(result)
            except Exception as e:
                print(f"任务{task_id}执行异常: {e}")
                results.append({"task_id": task_id, "status": "异常", "error": str(e)})
    
    save_all_validation_data()
    
    overall_end_time = datetime.now()
    overall_duration = (overall_end_time - overall_start_time).total_seconds()
    
    print("\n" + "=" * 60)
    print("所有任务执行完成!")
    print("=" * 60)
    
    success_count = sum(1 for r in results if r.get("status") == "成功")
    fail_count = num_tasks - success_count
    total_validations = sum(r.get("validation_count", 0) for r in results)
    
    print(f"\n执行统计:")
    print(f"  总任务数: {num_tasks}")
    print(f"  成功: {success_count}")
    print(f"  失败: {fail_count}")
    print(f"  总耗时: {overall_duration:.2f}秒")
    print(f"  平均耗时: {overall_duration/num_tasks:.2f}秒/任务")
    print(f"  字段校验测试总数: {total_validations}")
    print(f"  实际保存校验数据: {len(validation_data_list)}")
    
    print(f"\n注册数据已保存到: {EXCEL_FILE}")
    print(f"校验结果已保存到: {VALIDATION_EXCEL_FILE}")
    
    return results

if __name__ == "__main__":
    results = run_parallel_register(5)
